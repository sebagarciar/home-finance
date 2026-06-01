"""Santander España xlsx parser.

Layout (from sample):
    row 0..5 = bank header / metadata
    row 6    = column headers ('Transaction date', 'Value date', 'Description',
                                'Amount', 'Balance', 'Currency')
    row 7+   = data rows

Quirks:
    - Date format DD/MM/YYYY.
    - Amount uses U+2212 minus sign and European decimal: "−6,65".
    - txn_type is inferred from description (`PAGO MOVIL EN` -> card_payment,
      `COMPRA` -> card_payment, otherwise other / transfer / etc.).
"""
from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from typing import BinaryIO

from openpyxl import load_workbook

from ...models.transactions import TxnType
from .base import ParsedTxn, Parser

_AMOUNT_CLEAN = re.compile(r"[^\d,.\-−]")


def _parse_amount(raw: str | float | int) -> Decimal:
    if isinstance(raw, (int, float)):
        return Decimal(str(raw))
    s = str(raw).strip().replace("−", "-")  # U+2212 -> ASCII
    s = _AMOUNT_CLEAN.sub("", s)
    # European: "." thousands, "," decimal -> drop dots, swap comma to dot.
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    return Decimal(s) if s else Decimal("0")


def _infer_type(description: str, amount: Decimal) -> TxnType:
    d = description.lower()
    if d.startswith("pago movil en") or d.startswith("compra"):
        return TxnType.card_payment
    if "transferencia" in d or "transfer" in d or "bizum" in d:
        return TxnType.transfer
    if "abono" in d or d.startswith("ingreso"):
        return TxnType.deposit
    if "comision" in d:
        return TxnType.fee
    if "domiciliacion" in d or "recibo" in d:
        return TxnType.direct_debit
    return TxnType.other


class SantanderEsParser(Parser):
    name = "santander_es"

    def parse(self, file: BinaryIO) -> list[ParsedTxn]:
        wb = load_workbook(file, data_only=True)
        ws = wb[wb.sheetnames[0]]

        # Locate the header row by finding the row that contains "Transaction date".
        header_row_idx = None
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row and any(isinstance(c, str) and c.strip().lower() == "transaction date" for c in row):
                header_row_idx = i
                break
        if header_row_idx is None:
            raise ValueError("Could not locate header row in Santander ES xlsx")

        out: list[ParsedTxn] = []
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or all(c in (None, "") for c in row):
                continue
            txn_date_raw, _value_date, description, amount_raw, _balance, currency = row[:6]
            if not txn_date_raw or not description:
                continue
            if isinstance(txn_date_raw, datetime):
                txn_date = txn_date_raw.date()
            else:
                txn_date = datetime.strptime(str(txn_date_raw).strip(), "%d/%m/%Y").date()
            amount = _parse_amount(amount_raw)
            currency = (str(currency).strip() or "EUR").upper()
            descr = str(description).strip()
            out.append(
                ParsedTxn(
                    date=txn_date,
                    amount=amount,
                    currency=currency,
                    raw_description=descr,
                    txn_type=_infer_type(descr, amount),
                )
            )
        return out
